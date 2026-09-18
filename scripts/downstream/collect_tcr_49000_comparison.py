"""Collect audited 49000 results and provenance-separated TCR comparisons.

Activate pllm then protenix_abtcr; run from the repository root:
python scripts/downstream/collect_tcr_49000_comparison.py --output /new/report.json
Add --require-complete for final acceptance. Reads existing artifacts only; does
not infer, train, submit jobs, change original scores, or modify RESULTS.md.
Historical 29k controls are separate from the current-checkpoint task table.
--wait-for-complete can run under tmux to collect after the existing monitor
finishes; it fails visibly if that monitor exits or requests attention.
"""
from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path
import socket
import time

from scripts.downstream.monitor_tcr_49000_local import atomic, verify_report
from scripts.downstream.run_tcr_local_short import ROOT, now, sha

RUN = ROOT / 'output/downstream_generation/tcr_v5_49000_20260913'
OUTPUTS = ROOT / 'downstream/benchmark/outputs'
TAG = 'ours_tcr_v5_49000_20260913_runtime_v2'
EVALS = ('seen_test', 'seen_independent', 'unseen_independent')
OLD_TAGS = {
    'others_cdr3b': 'ours_v5_29000_others_cdr3b_masked_runtime_v2_nonspot_20260913',
    'cdr3ab': 'ours_v5_29000_others_cdr3ab_masked_runtime_v2_nonspot_20260913',
    'others_longab': 'ours_v5_29000_others_longab_masked_runtime_v1_20260913',
}
WORKBOOK = ROOT / ('downstream/benchmark/baselines/nm2025_supplement_20260914/'
                   '41592_2025_2910_MOESM10_ESM.xlsx')
WORKBOOK_SHA = '5e32c96d9a8a902d1abbec6a1b1110926e847d8f6c32178af274a2011431fad8'


class Sources:
    def __init__(self):
        self.hashes = {}

    def record(self, path):
        path = Path(path).resolve()
        self.hashes[str(path)] = sha(path)
        return path

    def json(self, path):
        return json.loads(self.record(path).read_text())

    def csv(self, path):
        with self.record(path).open(newline='') as handle:
            return list(csv.DictReader(handle))


def prediction_identity(rows):
    """Keep all source columns and ordering; only the predicted score may differ."""
    return [{k: v for k, v in row.items() if k != 'score'} for row in rows]


def protocol_comparison(current, previous):
    a, b = current['feature_manifest'], previous['feature_manifest']
    pa, pb = a['input_protocol'], b['input_protocol']
    semantic = lambda p: {k: v for k, v in p.items() if k != 'code_sha256'}
    changed = sorted(k for k in pa['code_sha256'].keys() | pb['code_sha256'].keys()
                     if pa['code_sha256'].get(k) != pb['code_sha256'].get(k))
    checks = {
        'head_config_equal': current['head_config'] == previous['head_config'],
        'input_semantics_equal': semantic(pa) == semantic(pb),
        'unique_input_hash_equal': a['pairs_sha256'] == b['pairs_sha256'],
        'archive_hash_equal': a['data_archive_md5'] == b['data_archive_md5'],
        'folds_equal': current['folds_requested'] == previous['folds_requested'],
        'eval_sets_equal': current['eval_sets_requested'] == previous['eval_sets_requested'],
    }
    if not all(checks.values()):
        raise ValueError(f'29k historical comparison protocol mismatch: {checks}')
    return {**checks, 'changed_code_files': changed,
            'strict_weight_only_causal_comparison': not changed,
            'interpretation': 'historical control, not a new rerun with fully frozen code'}


def old_binding(sources, current_tasks):
    rows = []
    base = OUTPUTS / 'tcr_binding_nm2025_retrained'
    for track, tag in OLD_TAGS.items():
        new = Path(current_tasks['t1-' + track]['result_dir'])
        old = base / tag / track / 'AS'
        nc, oc = [sources.json(p / 'run_config.json') for p in (new, old)]
        checks = protocol_comparison(nc, oc)
        ns, os = [sources.json(p / 'summary.json') for p in (new, old)]
        evidence = []
        for fold in range(1, 6):
            for split in EVALS:
                relative = Path(f'fold_{fold}') / split / 'predictions.csv'
                nr, oldr = [sources.csv(p / relative) for p in (new, old)]
                if prediction_identity(nr) != prediction_identity(oldr):
                    raise ValueError(f'29k/49k prediction row/label mismatch: {track}/{relative}')
                evidence.append({'fold': fold, 'split': split, 'rows': len(nr),
                                 'positives': sum(int(row['label']) for row in nr)})
        for split in EVALS:
            n, o = [s['eval_sets'][split] for s in (ns, os)]
            assert n['folds_complete'] == o['folds_complete'] == [1, 2, 3, 4, 5]
            values = {metric: {'ckpt29000': o['local'][metric + '_mean'],
                              'ckpt49000': n['local'][metric + '_mean'],
                              'delta_49k_minus_29k': n['local'][metric + '_mean'] - o['local'][metric + '_mean']}
                      for metric in ('auroc', 'auprc_precrec')}
            rows.append({'track': track, 'split': split, 'metrics': values,
                         'comparison_checks': checks, 'row_checks': evidence,
                         'old_checkpoint_sha256': oc['checkpoint_sha256']})
    return rows


def paper_binding(sources):
    from openpyxl import load_workbook
    assert sha(sources.record(WORKBOOK)) == WORKBOOK_SHA, 'publisher workbook drift'
    result = []
    mapping = [(0, 'large_beta', EVALS[0]), (1, 'large_beta', EVALS[1]),
               (2, 'large_beta', EVALS[2]), (3, 'others_multifeature', EVALS[0]),
               (4, 'others_beta', EVALS[0]), (5, 'others_multifeature', EVALS[1]),
               (6, 'others_beta', EVALS[1]), (7, 'others_multifeature', EVALS[2]),
               (8, 'others_beta', EVALS[2])]
    book = load_workbook(WORKBOOK, read_only=True, data_only=True)
    try:
        for index, scope, split in mapping:
            sheet, negative = book.worksheets[index], None
            for row_no, row in enumerate(sheet.iter_rows(min_row=3, values_only=True), 3):
                negative = row[0] or negative
                if negative != 'AS' or not row[1]:
                    continue
                counts = dict(zip(('TP', 'FP', 'TN', 'FN'), row[18:22]))
                n = sum(counts.values()) if all(isinstance(x, (int, float)) for x in counts.values()) else None
                result.append({'method': row[1], 'evidence': '[P]', 'data_scope': scope,
                               'split': split, 'sheet': sheet.title, 'row': row_no,
                               'auroc': row[2], 'auroc_std': row[3],
                               'auprc': row[4], 'auprc_std': row[5],
                               'reported_confusion_counts': counts, 'reported_n': n,
                               'reported_positive_fraction': (counts['TP'] + counts['FN']) / n if n else None,
                               'exact_local_rowset_match_established': False})
    finally:
        book.close()
    return result


def local_baselines(sources):
    groups = {}
    for group, filename, names in (
        ('tcr_clustering_embed', 'metrics.json', ('sceptr', 'tcrbert', 'esm2_150m', 'protbert', 'kmer')),
        ('tcr_representation', 'fewshot.json', ('sceptr', 'tcrdist', 'levenshtein', 'kmer', 'esm2_150m', 'protbert', 'tcrbert')),
        ('tcr_representation_paper6', 'fewshot.json', ('sceptr', 'tcrdist', 'levenshtein', 'kmer', 'esm2_150m', 'protbert', 'tcrbert')),
    ):
        groups[group] = {}
        for name in names:
            data = sources.json(OUTPUTS / group / name / filename)
            keys = ('protocol', 'shots', 'seeds', 'auroc_by_shot', 'input_fields', 'input_fields_note',
                    'n_units', 'n_epitopes', 'n_universe', 'n_background', 'n_train', 'n_test',
                    'k_sweep', 'seed', 'load_info', 'kmeans_ari_mean', 'kmeans_nmi_mean', 'kmeans_purity_mean')
            groups[group][name] = {k: data[k] for k in keys if k in data}
            groups[group][name]['evidence'] = '[L] historical local baseline; not re-executed'
            groups[group][name]['exact_49000_protocol_rerun'] = False
    groups['t2a_paper_and_published_artifact_calibration'] = sources.csv(OUTPUTS / 'tcr_clustering/calibration_report.csv')
    return groups


def generation_baselines(sources):
    result = {}
    base = OUTPUTS / 'tcr_generation_bench/setting_B'
    for name in ('tcrt5_official', 'gratcr_official', 'er_official', 'tcrdesign', 'tcrdiff', 'tcrdiff_held20', 'olga'):
        data = sources.json(base / name / 'metrics.json')
        result[name] = {view: {k: m[k] for k in ('d_edit', 'seq_recovery', 'f1', 'n_pmhc', 'n_pmhc_scored') if k in m}
                        for view, m in data['summary'].items()
                        if view in ('paper_sparse13', 'bioseq_unseen_common', 'held20', 'benchmark14')}
        result[name]['evidence'] = '[A] official stored predictions' if name.endswith('_official') else (
            '[C] unconditional control' if name == 'olga' else '[R] native-weight local pipeline')
        result[name]['same_information_and_sampling_established'] = False
    return result


def paper_representation(sources):
    groups = {}
    for row in sources.csv(OUTPUTS / 'external/paper_reported_baselines.csv'):
        if row['task'] != 'sceptr_six_pmhc_k200' or row['metric'] != 'auroc':
            continue
        # Only the display identity is corrected; the historical CSV is intact.
        name = 'ESM2-T6-8M' if row['method'] == 'ESM2-150M' else row['method']
        groups.setdefault(name, {})[row['split']] = float(row['value'])
    assert len(groups) == 6 and all(len(values) == 6 for values in groups.values())
    return {'source': 'https://arxiv.org/html/2406.06397v2',
            'source_location': 'Table SI, k=200; rechecked 2026-09-14',
            'evidence': '[P] mean derived from six printed three-decimal AUROCs',
            'exact_local_protocol_match': False,
            'legacy_ESM2_name_correction': 'ESM2-150M registry label actually refers to paper ESM2-T6-8M',
            'methods': {name: {'per_epitope': values, 'macro': sum(values.values()) / 6}
                        for name, values in groups.items()}}


def collect(require_complete=False):
    sources = Sources()
    plan = sources.json(RUN / 'monitor_plan_v3.json')
    latest = sources.json(RUN / 'monitor_v3/results_49000.json')
    assert latest['global_step'] == 49000 and latest['checkpoint_sha256'] == plan['checkpoint_sha256']
    verified = set()
    for group in plan['audit_groups']:
        if set(group['jobs']) & latest['tasks'].keys():
            verify_report(group)
            sources.record(group['report'])
            verified.update(group['jobs'])
    assert set(latest['tasks']) <= verified
    missing = sorted(set(plan['job_keys']) - latest['tasks'].keys())
    assert latest['complete'] == (not missing)
    if require_complete and missing:
        raise RuntimeError(f'not complete; unaudited tasks: {missing}')
    result = {'updated_utc': now(), 'complete': not missing, 'missing_tasks': missing,
              'current_49000': latest, 'old_checkpoint_comparison': old_binding(sources, latest['tasks']),
              'nm2025_paper_AS': paper_binding(sources), 'local_historical_baselines': local_baselines(sources),
              'generation_baselines': generation_baselines(sources),
              'sceptr_paper_baselines': paper_representation(sources),
              'old_checkpoint_missing_tasks': ['t1-cdr3b', 't2a', 't2b', 't3broad', 't3deep', 't4-benchmark14', 't4-held20'],
              'limitations': ['No new baseline or 29k inference in this collection.',
                             'Paper scopes, historical local runs, and current audited runs are distinct.',
                             'T3 broad historical scores precede grouping/multilabel scoring correction.',
                             'T3 paper ESM2 is T6 8M, NOT the local ESM2-150M baseline.',
                             'T4 common-6 aligns reference targets, not all model inputs or length priors.',
                             'Audit pass does not establish absent pretraining contamination.']}
    result['source_sha256'] = sources.hashes
    result['collector_sha256'] = sha(__file__)
    return result


def wait_for_complete(output):
    """Wait on the existing monitor; never create/restart an evaluation job."""
    status_path = output.with_suffix('.status.json')
    output.parent.mkdir(parents=True, exist_ok=True)
    if output.exists():
        raise FileExistsError(output)
    state = {'status': 'waiting_for_existing_monitor', 'start_utc': now(),
             'collector_sha256': sha(__file__), 'output': str(output.resolve())}
    # Reserving the status filename also prevents duplicate waiters for this output.
    with status_path.open('x') as handle:
        json.dump(state, handle, indent=2)
    try:
        while True:
            monitor = json.loads((RUN / 'monitor_v3/status.json').read_text())
            assert monitor['plan_sha256'] == sha(RUN / 'monitor_plan_v3.json')
            assert monitor['monitor_sha256'] == sha(ROOT / 'scripts/downstream/monitor_tcr_49000_local.py')
            state.update(updated_utc=now(), audited=len(monitor['audited']), monitor_status=monitor['status'])
            atomic(status_path, state)
            if monitor['status'] == 'all_tasks_executed_and_audits_passed':
                return state, status_path
            if monitor['status'] != 'running':
                raise RuntimeError(f'monitor needs attention: {monitor.get("error", monitor["status"])}')
            assert monitor['host'] == socket.gethostname()
            cmdline = Path(f'/proc/{monitor["pid"]}/cmdline').read_bytes().decode().replace('\0', ' ')
            assert monitor['plan'] in cmdline and 'monitor_tcr_49000_local.py' in cmdline, 'monitor PID disappeared/reused'
            print(now(), 'waiting; audited', state['audited'], flush=True)
            time.sleep(30)
    except Exception as exc:
        state.update(status='attention_required', error=repr(exc), updated_utc=now())
        atomic(status_path, state)
        raise


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--require-complete', action='store_true')
    parser.add_argument('--wait-for-complete', action='store_true',
                        help='wait for the existing healthy monitor, then require all ten audits')
    args = parser.parse_args()
    state, status_path = wait_for_complete(args.output) if args.wait_for_complete else (None, None)
    try:
        report = collect(args.require_complete or args.wait_for_complete)
        encoded = json.dumps(report, indent=2, ensure_ascii=False, allow_nan=False)
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with args.output.open('x') as handle:
            handle.write(encoded + '\n')
        if state is not None:
            state.update(status='complete', updated_utc=now(), output_sha256=sha(args.output))
            atomic(status_path, state)
    except Exception as exc:
        if state is not None:
            state.update(status='attention_required', updated_utc=now(), error=repr(exc))
            atomic(status_path, state)
        raise
    print(json.dumps({'output': str(args.output), 'complete': report['complete'], 'missing': report['missing_tasks']}))


if __name__ == '__main__':
    main()
